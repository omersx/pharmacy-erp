import io
import csv
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

COLUMNS = [
    {"header": "name_en *", "field": "name_en", "desc": "English name of the medicine (Required)"},
    {"header": "name_ar", "field": "name_ar", "desc": "Arabic name"},
    {"header": "sku", "field": "sku", "desc": "Stock Keeping Unit (Auto-generated if empty)"},
    {"header": "barcode", "field": "barcode", "desc": "Barcode for scanning"},
    {"header": "selling_price *", "field": "selling_price", "desc": "Selling price in base currency (Required)"},
    {"header": "base_unit *", "field": "base_unit", "desc": "Base unit (Required)"},
    {"header": "category", "field": "category", "desc": "Category name (Matched with existing)"},
    {"header": "generic_name", "field": "generic_name", "desc": "Generic name / Scientific name"},
    {"header": "brand_name", "field": "brand_name", "desc": "Brand name"},
    {"header": "strength", "field": "strength", "desc": "Strength (e.g., 500mg)"},
    {"header": "dosage_form", "field": "dosage_form", "desc": "Form of the medicine"},
    {"header": "manufacturer", "field": "manufacturer", "desc": "Manufacturer name"},
    {"header": "description", "field": "description", "desc": "Additional details"},
    {"header": "units_per_pack", "field": "units_per_pack", "desc": "Number of units per pack (Default: 1)"},
    {"header": "max_stock", "field": "max_stock", "desc": "Maximum stock allowed (Default: 0)"},
    {"header": "initial_stock", "field": "initial_stock", "desc": "Initial stock quantity"},
    {"header": "purchase_price", "field": "purchase_price", "desc": "Purchase price"},
    {"header": "batch_number", "field": "batch_number", "desc": "Batch/Lot number"},
    {"header": "expiry_date", "field": "expiry_date", "desc": "Expiry date YYYY-MM-DD"},
]

EXAMPLE_ROW = [
    "Amoxicillin 500mg",
    "أموكسيسيلين 500 ملغ",
    "AMX-500",
    "1234567890123",
    "15.50",
    "Box",
    "Antibiotics",
    "Amoxicillin",
    "Amoxil",
    "500mg",
    "capsule",
    "GSK",
    "Broad-spectrum antibiotic",
    "20",
    "100",
    "100",
    "10.50",
    "B-AMX-500",
    "2027-12-31"
]

BASE_UNITS = ["Pack", "Tablet", "Bottle", "Box", "Tube", "Strip", "Vial", "Ampoule", "Piece"]
DOSAGE_FORMS = ["tablet", "capsule", "syrup", "suspension", "injection", "infusion", "cream", "ointment", "gel", "drops", "eye_drops", "ear_drops", "nasal_spray", "inhaler", "suppository", "powder"]

def generate_csv_template() -> io.BytesIO:
    output = io.StringIO()
    writer = csv.writer(output)
    
    headers = [col["header"] for col in COLUMNS]
    descriptions = [col["desc"] for col in COLUMNS]
    
    writer.writerow(headers)
    writer.writerow(EXAMPLE_ROW)
    writer.writerow(descriptions)
    
    bio = io.BytesIO(output.getvalue().encode('utf-8'))
    return bio

def generate_xlsx_template() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Medicines Import"
    
    headers = [col["header"] for col in COLUMNS]
    descriptions = [col["desc"] for col in COLUMNS]
    
    ws.append(headers)
    ws.append(EXAMPLE_ROW)
    ws.append(descriptions)
    
    # Styles
    header_fill = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    example_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    desc_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    for col_idx, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    for cell in ws[2]:
        cell.fill = example_fill
        
    for cell in ws[3]:
        cell.fill = desc_fill
        
    # Auto-fit columns
    for col_idx, col in enumerate(COLUMNS, 1):
        col_letter = get_column_letter(col_idx)
        max_length = max(len(str(col["header"])), len(str(col["desc"])), len(str(EXAMPLE_ROW[col_idx-1])))
        ws.column_dimensions[col_letter].width = max_length + 2
        
    # Data Validation
    # base_unit is column 6 (F)
    dv_base_unit = DataValidation(type="list", formula1=f'"{",".join(BASE_UNITS)}"', allow_blank=False)
    ws.add_data_validation(dv_base_unit)
    dv_base_unit.add(f"F4:F1000")
    
    # dosage_form is column 11 (K)
    dv_dosage_form = DataValidation(type="list", formula1=f'"{",".join(DOSAGE_FORMS)}"', allow_blank=True)
    ws.add_data_validation(dv_dosage_form)
    dv_dosage_form.add(f"K4:K1000")
    
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
