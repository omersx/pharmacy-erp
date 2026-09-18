from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from app.core.audit import log_action
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, or_
from app.core.database import get_db
from app.core.deps import get_current_user
from app.modules.auth.models import User
from app.modules.medicines.models import Medicine, MedicineBatch, MedicineCategory
from app.modules.medicines.schemas import (
    MedicineCreate, MedicineUpdate, MedicineResponse, 
    BatchCreate, BatchUpdate, BatchResponse, BatchConflictResponse,
    CategoryCreate, CategoryResponse,
    CategoryUpdate, CategoryTreeResponse, MedicineCreateWithBatch, MedicineDetailResponse,
    BulkImportResponse, BulkImportRowError
)
from app.modules.inventory.models import InventoryMovement
from app.modules.organizations.models import Branch
from typing import List
from uuid import UUID
from datetime import datetime
from fastapi.responses import JSONResponse, StreamingResponse
import csv
import io
from openpyxl import load_workbook
from app.modules.medicines.template_generator import generate_csv_template, generate_xlsx_template, BASE_UNITS, DOSAGE_FORMS

router = APIRouter(tags=["Medicines"])

@router.get("/", response_model=List[MedicineResponse])
async def list_medicines(db: AsyncSession = Depends(get_db)):
    res = await db.execute(select(Medicine).where(Medicine.is_active == True))
    return res.scalars().all()

@router.get("/search", response_model=List[MedicineResponse])
async def search_medicines(q: str, db: AsyncSession = Depends(get_db)):
    query = select(Medicine).where(
        Medicine.is_active == True,
        or_(
            Medicine.name_en.ilike(f"%{q}%"),
            Medicine.name_ar.ilike(f"%{q}%"),
            Medicine.barcode == q,
            Medicine.sku == q
        )
    ).limit(50)
    res = await db.execute(query)
    return res.scalars().all()

@router.post("/", response_model=MedicineResponse)
async def create_medicine(data: MedicineCreate, request: Request, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    medicine = Medicine(**data.model_dump())
    db.add(medicine)
    await db.commit()
    await db.refresh(medicine)
    
    await log_action(
        db=db, action="CREATE", module="Medicine",
        details=f"Created medicine: {medicine.name_en}",
        user=current_user, entity_type="Medicine", entity_id=str(medicine.id),
        ip_address=request.client.host if request.client else None
    )
    
    return medicine

@router.put("/{id}", response_model=MedicineResponse)
async def update_medicine(id: UUID, data: MedicineUpdate, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    medicine = await db.get(Medicine, id)
    if not medicine:
        raise HTTPException(status_code=404, detail="Medicine not found")
        
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(medicine, key, value)
        
    await db.commit()
    await db.refresh(medicine)
    return medicine

@router.delete("/{id}")
async def delete_medicine(id: UUID, request: Request, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    medicine = await db.get(Medicine, id)
    if not medicine:
        raise HTTPException(status_code=404, detail="Medicine not found")
        
    medicine.is_active = False
    await db.commit()
    
    await log_action(
        db=db, action="DELETE", module="Medicine",
        details=f"Deleted medicine: {medicine.name_en}",
        user=current_user, entity_type="Medicine", entity_id=str(medicine.id),
        ip_address=request.client.host if request.client else None
    )
    
    return {"message": "Medicine deleted successfully"}

@router.get("/{id}/batches", response_model=List[BatchResponse])
async def list_batches(id: UUID, db: AsyncSession = Depends(get_db)):
    res = await db.execute(
        select(MedicineBatch)
        .where(MedicineBatch.medicine_id == id)
        .order_by(MedicineBatch.expiry_date.asc())
    )
    return res.scalars().all()

@router.post("/{id}/batches", response_model=BatchResponse)
async def add_batch(id: UUID, data: BatchCreate, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Add a new batch with smart duplicate detection.
    
    - Same batch number + same expiry → merge quantity into existing batch
    - Same batch number + different expiry → return 409 conflict warning
    - New batch number → create new batch
    """
    medicine = await db.get(Medicine, id)
    if not medicine:
        raise HTTPException(status_code=404, detail="Medicine not found")
    
    # Check for existing batch with same batch_number
    existing_res = await db.execute(
        select(MedicineBatch).where(
            MedicineBatch.medicine_id == id,
            MedicineBatch.batch_number == data.batch_number
        )
    )
    existing_batch = existing_res.scalars().first()
    
    if existing_batch:
        # Same batch number exists - check expiry date
        existing_expiry = existing_batch.expiry_date.date() if existing_batch.expiry_date else None
        new_expiry = data.expiry_date.date() if data.expiry_date else None
        
        if existing_expiry == new_expiry:
            # MERGE: Same batch + same expiry → add quantity to existing
            existing_batch.quantity_received = float(existing_batch.quantity_received) + float(data.quantity_received)
            existing_batch.quantity_remaining = float(existing_batch.quantity_remaining) + float(data.quantity_received)
            existing_batch.is_active = True
            if existing_batch.status == "OUT_OF_STOCK":
                existing_batch.status = "ACTIVE"
            
            # Get branch for movement
            branch_res = await db.execute(select(Branch).limit(1))
            branch = branch_res.scalars().first()
            
            # Log stock receipt movement
            movement = InventoryMovement(
                medicine_id=id,
                batch_id=existing_batch.id,
                branch_id=branch.id if branch else None,
                movement_type="STOCK_RECEIPT",
                quantity=float(data.quantity_received),
                reference_type="BATCH_MERGE",
                notes=f"Merged {data.quantity_received} units into existing batch {data.batch_number}",
                created_by=current_user.id
            )
            db.add(movement)
            
            await db.commit()
            await db.refresh(existing_batch)
            return existing_batch
        else:
            # CONFLICT: Same batch number but different expiry
            return JSONResponse(
                status_code=409,
                content={
                    "conflict": True,
                    "message": f"Batch {data.batch_number} already exists with expiry {existing_expiry.isoformat() if existing_expiry else 'unknown'}. The entered expiry date is {new_expiry.isoformat() if new_expiry else 'unknown'}. Please verify.",
                    "existing_batch_number": data.batch_number,
                    "existing_expiry_date": existing_expiry.isoformat() if existing_expiry else None,
                    "entered_expiry_date": new_expiry.isoformat() if new_expiry else None
                }
            )
    
    # NEW BATCH: No duplicate found
    batch = MedicineBatch(
        medicine_id=id,
        quantity_remaining=data.quantity_received,
        status="ACTIVE",
        **data.model_dump()
    )
    db.add(batch)
    await db.flush()
    
    # Get branch for movement
    branch_res = await db.execute(select(Branch).limit(1))
    branch = branch_res.scalars().first()
    
    # Log stock receipt movement
    movement = InventoryMovement(
        medicine_id=id,
        batch_id=batch.id,
        branch_id=branch.id if branch else None,
        movement_type="STOCK_RECEIPT",
        quantity=float(data.quantity_received),
        reference_type="NEW_BATCH",
        notes=f"New batch {data.batch_number} added with {data.quantity_received} units",
        created_by=current_user.id
    )
    db.add(movement)
    
    await db.commit()
    await db.refresh(batch)
    return batch

@router.put("/{id}/batches/{batch_id}", response_model=BatchResponse)
async def update_batch(
    id: UUID, batch_id: UUID, data: BatchUpdate, 
    db: AsyncSession = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    """Edit batch details. Quantity cannot be changed directly — use stock adjustments."""
    batch = await db.get(MedicineBatch, batch_id)
    if not batch or batch.medicine_id != id:
        raise HTTPException(status_code=404, detail="Batch not found")
    
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(batch, key, value)
    
    await db.commit()
    await db.refresh(batch)
    return batch

@router.delete("/{id}/batches/{batch_id}")
async def deactivate_batch(
    id: UUID, batch_id: UUID, 
    db: AsyncSession = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    """Soft-deactivate a batch. It will no longer be available for sales."""
    batch = await db.get(MedicineBatch, batch_id)
    if not batch or batch.medicine_id != id:
        raise HTTPException(status_code=404, detail="Batch not found")
    
    batch.is_active = False
    batch.status = "DEACTIVATED"
    await db.commit()
    return {"message": f"Batch {batch.batch_number} deactivated"}

@router.get("/categories", response_model=List[CategoryResponse])
async def list_categories(db: AsyncSession = Depends(get_db)):
    res = await db.execute(select(MedicineCategory))
    return res.scalars().all()

@router.post("/categories", response_model=CategoryResponse)
async def create_category(data: CategoryCreate, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    category = MedicineCategory(**data.model_dump())
    db.add(category)
    await db.commit()
    await db.refresh(category)
    return category

@router.get("/categories/tree")
async def get_category_tree(db: AsyncSession = Depends(get_db)):
    """Return categories as a nested tree structure"""
    res = await db.execute(select(MedicineCategory).order_by(MedicineCategory.sort_order))
    all_cats = res.scalars().all()
    
    # Build tree: top-level (parent_id is None) with children nested
    cat_map = {}
    for cat in all_cats:
        cat_dict = {
            "id": str(cat.id),
            "name_en": cat.name_en,
            "name_ar": cat.name_ar,
            "parent_id": str(cat.parent_id) if cat.parent_id else None,
            "category_level": cat.category_level,
            "icon": cat.icon,
            "sort_order": cat.sort_order,
            "children": []
        }
        cat_map[str(cat.id)] = cat_dict
    
    tree = []
    for cat_id, cat_dict in cat_map.items():
        if cat_dict["parent_id"] and cat_dict["parent_id"] in cat_map:
            cat_map[cat_dict["parent_id"]]["children"].append(cat_dict)
        else:
            tree.append(cat_dict)
    
    return tree

@router.put("/categories/{id}", response_model=CategoryResponse)
async def update_category(id: UUID, data: CategoryUpdate, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    cat = await db.get(MedicineCategory, id)
    if not cat:
        raise HTTPException(status_code=404, detail="Category not found")
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(cat, key, value)
    await db.commit()
    await db.refresh(cat)
    return cat

@router.delete("/categories/{id}")
async def delete_category(id: UUID, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    # Check if any medicines use this category
    res = await db.execute(select(Medicine).where(Medicine.category_id == id).limit(1))
    if res.scalars().first():
        raise HTTPException(status_code=400, detail="Cannot delete category with assigned medicines")
    cat = await db.get(MedicineCategory, id)
    if not cat:
        raise HTTPException(status_code=404, detail="Category not found")
    await db.delete(cat)
    await db.commit()
    return {"message": "Category deleted"}

@router.get("/import-template")
async def get_import_template(format: str = "csv"):
    if format == "xlsx":
        bio = generate_xlsx_template()
        return StreamingResponse(
            bio,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=medicines_import_template.xlsx"}
        )
    else:
        bio = generate_csv_template()
        return StreamingResponse(
            bio,
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=medicines_import_template.csv"}
        )

@router.post("/bulk-import", response_model=BulkImportResponse)
async def bulk_import_medicines(
    request: Request,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    content = await file.read()
    filename = file.filename.lower()
    
    rows = []
    
    if filename.endswith('.csv'):
        text_content = content.decode('utf-8')
        reader = csv.DictReader(io.StringIO(text_content))
        all_rows = list(reader)
        for i, row in enumerate(all_rows):
            name_en_val = str(row.get('name_en *') or '')
            if i < 2 and (name_en_val == 'Amoxicillin 500mg' or 'Required' in name_en_val):
                continue
            mapped_row = {
                'name_en': row.get('name_en *'),
                'name_ar': row.get('name_ar'),
                'sku': row.get('sku'),
                'barcode': row.get('barcode'),
                'selling_price': row.get('selling_price *'),
                'base_unit': row.get('base_unit *'),
                'category': row.get('category'),
                'generic_name': row.get('generic_name'),
                'brand_name': row.get('brand_name'),
                'strength': row.get('strength'),
                'dosage_form': row.get('dosage_form'),
                'manufacturer': row.get('manufacturer'),
                'description': row.get('description'),
                'units_per_pack': row.get('units_per_pack'),
                'reorder_level': row.get('reorder_level'),
                'max_stock': row.get('max_stock'),
                'requires_prescription': row.get('requires_prescription'),
                'is_controlled': row.get('is_controlled'),
                'is_cold_chain': row.get('is_cold_chain'),
                'allow_loose_sale': row.get('allow_loose_sale'),
                'initial_stock': row.get('initial_stock'),
                'purchase_price': row.get('purchase_price'),
                'batch_number': row.get('batch_number'),
                'expiry_date': row.get('expiry_date')
            }
            rows.append((i+2, mapped_row))
    elif filename.endswith('.xlsx'):
        wb = load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None for v in row):
                continue
            
            row_dict = dict(zip(headers, row))
            name_en_val = str(row_dict.get('name_en *') or '')
            if row_idx <= 3 and (name_en_val == 'Amoxicillin 500mg' or 'Required' in name_en_val):
                continue
                
            mapped_row = {
                'name_en': row_dict.get('name_en *'),
                'name_ar': row_dict.get('name_ar'),
                'sku': row_dict.get('sku'),
                'barcode': row_dict.get('barcode'),
                'selling_price': row_dict.get('selling_price *'),
                'base_unit': row_dict.get('base_unit *'),
                'category': row_dict.get('category'),
                'generic_name': row_dict.get('generic_name'),
                'brand_name': row_dict.get('brand_name'),
                'strength': row_dict.get('strength'),
                'dosage_form': row_dict.get('dosage_form'),
                'manufacturer': row_dict.get('manufacturer'),
                'description': row_dict.get('description'),
                'units_per_pack': row_dict.get('units_per_pack'),
                'reorder_level': row_dict.get('reorder_level'),
                'max_stock': row_dict.get('max_stock'),
                'requires_prescription': row_dict.get('requires_prescription'),
                'is_controlled': row_dict.get('is_controlled'),
                'is_cold_chain': row_dict.get('is_cold_chain'),
                'allow_loose_sale': row_dict.get('allow_loose_sale'),
                'initial_stock': row_dict.get('initial_stock'),
                'purchase_price': row_dict.get('purchase_price'),
                'batch_number': row_dict.get('batch_number'),
                'expiry_date': row_dict.get('expiry_date')
            }
            rows.append((row_idx, mapped_row))
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format")

    errors = []
    medicines_to_create = []
    valid_rows = []
    
    cat_res = await db.execute(select(MedicineCategory))
    categories = {c.name_en.lower(): c.id for c in cat_res.scalars().all()}
    
    sku_res = await db.execute(select(Medicine.sku))
    existing_skus = set(sku_res.scalars().all())
    
    barcode_res = await db.execute(select(Medicine.barcode).where(Medicine.barcode.isnot(None)))
    existing_barcodes = set(barcode_res.scalars().all())
    
    file_skus = set()
    file_barcodes = set()
    
    import time
    timestamp = int(time.time())
    
    for row_idx, row in rows:
        row_errors = []
        
        name_en = row.get('name_en')
        if not name_en:
            row_errors.append(BulkImportRowError(row=row_idx, field="name_en", message="name_en is required"))
            
        try:
            selling_price = float(row.get('selling_price') or 0)
            if selling_price <= 0:
                row_errors.append(BulkImportRowError(row=row_idx, field="selling_price", message="selling_price must be > 0"))
        except ValueError:
            row_errors.append(BulkImportRowError(row=row_idx, field="selling_price", message="Invalid selling_price"))
            
        base_unit = row.get('base_unit')
        if not base_unit or base_unit not in BASE_UNITS:
            row_errors.append(BulkImportRowError(row=row_idx, field="base_unit", message=f"base_unit must be one of {BASE_UNITS}"))
            
        dosage_form = row.get('dosage_form')
        if dosage_form and dosage_form not in DOSAGE_FORMS:
            row_errors.append(BulkImportRowError(row=row_idx, field="dosage_form", message=f"dosage_form must be one of {DOSAGE_FORMS}"))
            
        sku = str(row.get('sku') or "").strip()
        if not sku:
            sku = f"IMP-{timestamp}-{row_idx}"
            
        if sku in existing_skus or sku in file_skus:
            row_errors.append(BulkImportRowError(row=row_idx, field="sku", message=f"Duplicate SKU: {sku}"))
        else:
            file_skus.add(sku)
            
        barcode = str(row.get('barcode') or "").strip()
        if barcode:
            if barcode in existing_barcodes or barcode in file_barcodes:
                row_errors.append(BulkImportRowError(row=row_idx, field="barcode", message=f"Duplicate barcode: {barcode}"))
            else:
                file_barcodes.add(barcode)
        else:
            barcode = None
            
        category_id = None
        category_name = str(row.get('category') or "").strip()
        if category_name:
            cat_lower = category_name.lower()
            if cat_lower in categories:
                category_id = categories[cat_lower]
            else:
                new_cat = MedicineCategory(name_en=category_name)
                db.add(new_cat)
                await db.flush()
                categories[cat_lower] = new_cat.id
                category_id = new_cat.id
                
        def parse_bool(val):
            if not val: return False
            return str(val).strip().lower() in ['yes', 'true', '1']
            
        def parse_int(val, default):
            try:
                return int(val)
            except (ValueError, TypeError):
                return default
                
        if row_errors:
            errors.extend(row_errors)
        else:
            medicines_to_create.append(
                Medicine(
                    name_en=name_en,
                    name_ar=row.get('name_ar'),
                    sku=sku,
                    barcode=barcode,
                    selling_price=selling_price,
                    base_unit=base_unit,
                    category_id=category_id,
                    generic_name=row.get('generic_name'),
                    brand_name=row.get('brand_name'),
                    strength=row.get('strength'),
                    dosage_form=dosage_form,
                    manufacturer=row.get('manufacturer'),
                    description=row.get('description'),
                    units_per_pack=parse_int(row.get('units_per_pack'), 1),
                    reorder_level=parse_int(row.get('reorder_level'), 10),
                    max_stock=parse_int(row.get('max_stock'), 0),
                    requires_prescription=parse_bool(row.get('requires_prescription')),
                    is_controlled=parse_bool(row.get('is_controlled')),
                    is_cold_chain=parse_bool(row.get('is_cold_chain')),
                    allow_loose_sale=parse_bool(row.get('allow_loose_sale')),
                    is_active=True
                )
            )
            valid_rows.append((row_idx, row))
            
    if medicines_to_create:
        db.add_all(medicines_to_create)
        await db.commit()
        for m in medicines_to_create:
            await db.refresh(m)
            
        batches_to_create = []
        for m, (row_idx, row_data) in zip(medicines_to_create, valid_rows):
            stock = parse_int(row_data.get('initial_stock'), 0)
            if stock > 0:
                expiry_str = row_data.get('expiry_date')
                parsed_date = None
                if expiry_str:
                    try:
                        parsed_date = datetime.strptime(str(expiry_str).strip(), "%Y-%m-%d").date()
                    except ValueError:
                        pass
                
                try:
                    purchase_price = float(row_data.get('purchase_price') or 0)
                    if purchase_price <= 0:
                        purchase_price = float(m.selling_price * 0.7)
                except (ValueError, TypeError):
                    purchase_price = float(m.selling_price * 0.7)
                    
                batch = MedicineBatch(
                    medicine_id=m.id,
                    batch_number=str(row_data.get('batch_number') or "").strip() or f"B-{m.sku}",
                    purchase_price=purchase_price,
                    quantity_received=stock,
                    quantity_remaining=stock,
                    expiry_date=parsed_date,
                    received_date=datetime.now()
                )
                batches_to_create.append(batch)
                
        if batches_to_create:
            db.add_all(batches_to_create)
            await db.commit()
            
    await log_action(
        db=db, action="IMPORT", module="Medicine",
        details=f"Bulk imported {len(medicines_to_create)} medicines ({len(set([e.row for e in errors]))} failed)",
        user=current_user, ip_address=request.client.host if request.client else None
    )
            
    return BulkImportResponse(
        total_rows=len(rows),
        successful=len(medicines_to_create),
        failed=len(set([e.row for e in errors])),
        errors=errors,
        created_medicines=[MedicineResponse.model_validate(m) for m in medicines_to_create]
    )

@router.post("/with-batch", response_model=MedicineResponse)
async def create_medicine_with_batch(
    data: MedicineCreateWithBatch, 
    db: AsyncSession = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    """Create medicine and optionally add first batch in one transaction"""
    batch_data = data.initial_batch
    medicine_data = data.model_dump(exclude={"initial_batch"})
    
    medicine = Medicine(**medicine_data)
    db.add(medicine)
    await db.flush()  # get medicine.id without committing
    
    if batch_data:
        batch_number = batch_data.batch_number or f"BATCH-{medicine.sku}-001"
        batch = MedicineBatch(
            medicine_id=medicine.id,
            batch_number=batch_number,
            expiry_date=batch_data.expiry_date,
            purchase_price=batch_data.purchase_price,
            quantity_received=batch_data.quantity,
            quantity_remaining=batch_data.quantity,
            production_date=batch_data.production_date,
            supplier_id=batch_data.supplier_id
        )
        db.add(batch)
    
    await db.commit()
    await db.refresh(medicine)
    return medicine

@router.get("/{id}", response_model=MedicineDetailResponse)
async def get_medicine(id: UUID, db: AsyncSession = Depends(get_db)):
    medicine = await db.get(Medicine, id)
    if not medicine:
        raise HTTPException(status_code=404, detail="Medicine not found")
    
    # Get batches
    batch_res = await db.execute(
        select(MedicineBatch).where(MedicineBatch.medicine_id == id)
    )
    batches = batch_res.scalars().all()
    
    # Get category name
    category_name = None
    if medicine.category_id:
        cat = await db.get(MedicineCategory, medicine.category_id)
        if cat:
            category_name = cat.name_en
    
    return {
        **{c.name: getattr(medicine, c.name) for c in medicine.__table__.columns},
        "batches": batches,
        "category_name": category_name
    }
